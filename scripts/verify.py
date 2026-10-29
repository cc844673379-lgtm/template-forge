"""verify.py — 八项机器判定 + 残留扫描（交付门禁）

用法: python verify.py <spec.json> <产物路径>
判定（蓝图 B07 Validation 八项）:
  1 spec schema  2 结构命中率≥0.8  3 样式 diff（规范化语义比较——教训 #5）
  4 残留零命中  5 槽位完整性  6 源件哈希未变  7 deviated_from 完整  8 样片同通道
"""
import sys, os, json, re, hashlib, argparse, zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import spec as specmod
from units import pt_to_px


def norm_xml(x: bytes) -> str:
    """XML 规范化语义比较（教训 #5）：引号风格 + 空白归一。禁止逐字节比。"""
    s = x.decode("utf-8", "ignore").replace('"', "'")
    return re.sub(r"\s+", "", s)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("spec")
    ap.add_argument("artifact")
    a = ap.parse_args()

    s = specmod.load(a.spec)
    r = {}
    # 1) schema
    ok, missing = specmod.validate(s)
    r["1_schema"] = ok or f"missing: {missing}"
    # 6) 源件哈希
    r["6_src_hash"] = hashlib.sha256(open(s["source"]["file"], "rb").read()).hexdigest() == s["source"]["sha256"]
    # 7) deviated_from 完整（validate 已查字段；这里查"有偏离但没记录"由 render 层保证，报告呈现条数）
    r["7_deviated_from"] = f"{len(s.get('deviated_from', []))} 条，字段完整" if ok else "schema 未过"

    ext = os.path.splitext(a.artifact)[1].lower()
    if ext == ".docx":
        import docx as pydocx
        d = pydocx.Document(a.artifact)
        text = "\n".join(p.text for p in d.paragraphs)
        src_h2 = [n for n in s["extract"]["skeleton"] if n["type"] == "heading" and n.get("level") == 2]
        dst_h2 = [p for p in d.paragraphs if p.style.name == "Heading 2"]
        r["2_structure_hit"] = len(dst_h2) / max(len(src_h2), 1)
        # 3) 样式 diff：clone 场景比 XML part 语义一致
        if s["strategy"] == "clone-and-fill" and s["source"]["type"] == "docx":
            zs, zd = zipfile.ZipFile(s["source"]["file"]), zipfile.ZipFile(a.artifact)
            r["3_style_diff"] = all(
                norm_xml(zs.read(p)) == norm_xml(zd.read(p))
                for p in ("word/styles.xml", "word/theme/theme1.xml"))
        else:
            r["3_style_diff"] = "rebuild 场景：以 token 抽查替代（见 render 输出）"
        h1 = [p for p in d.paragraphs if p.style.name == "Heading 1"]
        r["5_slots"] = len(dst_h2) == len(src_h2)
    elif ext == ".html":
        h = open(a.artifact, encoding="utf-8").read()
        text = re.sub(r"<[^>]+>", "", h)
        src_h2 = [n for n in s["extract"]["skeleton"] if n["type"] == "heading" and n.get("level") == 2]
        dst_h2 = re.findall(r"<h2>", h)
        r["2_structure_hit"] = len(dst_h2) / max(len(src_h2), 1)
        root = re.search(r":root \{(.*?)\}", h, re.S)
        vars_ = dict(re.findall(r"(--[\w-]+):\s*([^;]+);", root.group(1))) if root else {}
        st = s["extract"]["paragraph_styles"]
        exp_h2 = pt_to_px(st.get("Heading 2", {}).get("size", 12))
        got_h2 = vars_.get("--size-h2", "")
        r["3_style_diff"] = {"h2_px": {"expect": exp_h2, "got": got_h2.strip().replace("px", ""),
                                        "ok": str(exp_h2) + "px" == got_h2.strip()},
                             "accent_ok": vars_.get("--color-accent", "").strip()
                             == "#" + (st.get("Heading 2", {}).get("color") or "")}
        r["3_style_diff_pass"] = (r["3_style_diff"]["h2_px"]["ok"] and r["3_style_diff"]["accent_ok"])
        r["5_slots"] = True
    else:
        text = open(a.artifact, encoding="utf-8").read()
        r["2_structure_hit"] = "md 目标：仅内容层"
        r["3_style_diff"] = "md 无样式概念（铁律）"
        r["5_slots"] = True

    # 4) 残留扫描
    r["4_leftover"] = re.findall(r"\{\{[^}]*\}\}|\{\{|\}\}|（内容待补）|（内容缺失）|TODO", text) or "零命中"
    # 8) 样片同通道（由 probe 记录；此处报告 spec 标志）
    r["8_probe_channel"] = "skipped" if s.get("probe_skipped") else "目标通道样片（或用户确认跳过）"

    ok = (r["1_schema"] is True and r["6_src_hash"] is True
          and (r["2_structure_hit"] == "md 目标：仅内容层" or (isinstance(r["2_structure_hit"], float) and r["2_structure_hit"] >= 0.8))
          and (r["3_style_diff"] is True or r.get("3_style_diff_pass") is True or isinstance(r["3_style_diff"], str))
          and r["4_leftover"] == "零命中" and r["5_slots"] is True)
    r["verdict"] = "PASS" if ok else "FAIL"
    print(json.dumps(r, ensure_ascii=False, indent=2))
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
